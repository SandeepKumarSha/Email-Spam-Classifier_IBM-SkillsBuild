import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set page margins (1 inch around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Color Palette Definitions
    PRIMARY_COLOR = RGBColor(30, 58, 138)      # Deep Navy #1E3A8A
    SECONDARY_COLOR = RGBColor(13, 148, 136)   # Teal #0D9488
    TEXT_COLOR = RGBColor(30, 41, 59)          # Charcoal #1E293B
    MUTED_COLOR = RGBColor(100, 116, 139)      # Slate Gray #64748B

    # Base normal style setup
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = TEXT_COLOR

    # --- TITLE PAGE / HEADER ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(6)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("EMAIL SPAM CLASSIFIER")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(20)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Project Implementation Report & UI Demonstration")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(16)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Title:", "Email Spam Classifier Web Application"),
        ("Author / Student:", "Sandeep Kumar Sha"),
        ("Program / Platform:", "IBM SkillsBuild Internship Program"),
        ("Domain:", "Natural Language Processing (NLP) & Machine Learning"),
        ("Technology Stack:", "Python, Flask, Scikit-Learn, NLTK, HTML5/CSS3")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        # Format label
        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.color.rgb = PRIMARY_COLOR
        cell_lbl.width = Inches(2.2)
        
        # Format value
        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.color.rgb = TEXT_COLOR
        cell_val.width = Inches(4.3)
        
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # Helper function for headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
        return h

    def add_body_p(text, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT_COLOR
        return p

    def add_bullet_p(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.color.rgb = PRIMARY_COLOR
        r_text = p.add_run(text)
        r_text.font.color.rgb = TEXT_COLOR
        return p

    def add_figure(image_path, caption):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(10)
            p_img.paragraph_format.space_after = Pt(4)
            run_img = p_img.add_run()
            run_img.add_picture(image_path, width=Inches(5.8))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Calibri'
            run_cap.font.size = Pt(9.5)
            run_cap.font.italic = True
            run_cap.font.color.rgb = MUTED_COLOR

    # --- 1. EXECUTIVE SUMMARY ---
    add_heading_1("1. Executive Summary")
    add_body_p("The Email Spam Classifier is an end-to-end Machine Learning and Natural Language Processing (NLP) web application built to analyze incoming messages and emails in real-time, categorizing them as either 'Spam' or 'Legitimate (Ham)'. Utilizing a Multinomial Naive Bayes classification model trained on TF-IDF vectorized text representations, the system achieves an overall accuracy of 97.09% and a precision score of 100.00%.")
    add_body_p("The project features a sleek, responsive glassmorphic user interface developed with Flask, HTML5, and CSS3. Users can paste custom email content or select interactive pre-configured message samples to evaluate the classifier. This report documents the complete system architecture, dataset pipeline, model performance, user interface designs, and real-time execution outputs.")

    # --- 2. PROBLEM STATEMENT & OBJECTIVES ---
    add_heading_1("2. Project Objectives & Scope")
    add_body_p("Electronic mail remains a primary channel for communication; however, the influx of spam, promotional junk, and fraudulent phishing messages poses significant cybersecurity risks and reduces user productivity.")
    add_heading_2("Key Objectives:")
    add_bullet_p("Automated Filtering: ", "Develop an accurate NLP pipeline to automatically distinguish unsolicited spam from valid emails.")
    add_bullet_p("Zero False Positives: ", "Optimize model precision to 100%, ensuring that legitimate messages are never misclassified as spam.")
    add_bullet_p("Interactive Web Interface: ", "Provide a modern, web-accessible user interface for instant message analysis with visual confidence scores.")
    add_bullet_p("Scalable Deployment: ", "Build a lightweight web architecture suitable for cloud hosting and scalable API integration.")

    # --- 3. SYSTEM ARCHITECTURE & TECH STACK ---
    add_heading_1("3. System Architecture & Tech Stack")
    add_body_p("The system follows a modular architecture separating data processing, model inference, and web presentation layers:")
    
    add_bullet_p("Machine Learning & NLP Framework: ", "Python 3, Scikit-Learn (MultinomialNB, TF-IDF Vectorizer), NLTK (PorterStemmer for word stemming).")
    add_bullet_p("Web Framework: ", "Flask (Python backend micro-framework for handling HTTP POST requests and model inference).")
    add_bullet_p("Frontend Design: ", "HTML5, CSS3, Google Fonts (Inter, Poppins), and FontAwesome icons.")
    add_bullet_p("Model Storage: ", "Serialized pickle artifacts (model.pkl and vectorizer.pkl) for fast startup loading.")

    # Pipeline block text
    p_pipe = doc.add_paragraph()
    p_pipe.paragraph_format.space_before = Pt(8)
    p_pipe.paragraph_format.space_after = Pt(12)
    p_pipe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pipe = p_pipe.add_run("Input Text → Lowercasing & Cleaning → Stopword & Punctuation Removal → Porter Stemming → TF-IDF Vectorization → Naive Bayes Classifier → Output Prediction & Confidence Score")
    r_pipe.font.bold = True
    r_pipe.font.size = Pt(10)
    r_pipe.font.color.rgb = SECONDARY_COLOR

    # --- 4. MODEL EVALUATION & PERFORMANCE ---
    add_heading_1("4. Model Performance & Evaluation")
    add_body_p("The Multinomial Naive Bayes model was evaluated against test dataset samples. The evaluation metrics and confusion matrix are highlighted below:")

    # Table of metrics
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric", "Score", "Description"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].paragraphs[0].add_run(title).font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=150, right=150)

    rows_data = [
        ("Accuracy", "97.09%", "Percentage of overall correctly classified messages across test dataset."),
        ("Precision", "100.00%", "Ratio of true spam detections out of all predicted spam (Zero False Positives achieved).")
    ]

    for idx, row in enumerate(rows_data):
        row_cells = table.rows[idx + 1].cells
        for c_idx, val in enumerate(row):
            p = row_cells[c_idx].paragraphs[0]
            r = p.add_run(val)
            if c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = SECONDARY_COLOR
            set_cell_background(row_cells[c_idx], "F8FAFC" if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Insert confusion matrix and evaluation score graphics
    add_heading_2("Confusion Matrix & Metrics Visualizations")
    add_figure("static/confusion_matrix.png", "Figure 1: Confusion Matrix of Spam Classifier Model")
    add_figure("static/evaluation_score.png", "Figure 2: Accuracy and Precision Metric Scores")

    # --- 5. USER INTERFACE & OUTPUT DEMONSTRATION ---
    add_heading_1("5. User Interface & Output Demonstration")
    add_body_p("The web application features a clean, responsive card interface called 'Spam Shield'. Below are actual captured screenshots demonstrating the initial UI layout and live classification outputs for both Spam and Legitimate (Ham) emails.")

    # UI Screenshot 1: Home Page
    add_heading_2("5.1 Application Home Page UI")
    add_body_p("The home interface presents a welcoming input section with quick example preset buttons allowing users to test sample messages with one click.")
    add_figure("static/doc_ui_home.png", "Figure 3: Application Interface - Main Input Form and Example Message Options")

    # UI Screenshot 2: Spam Output
    add_heading_2("5.2 Spam Detection Output Interface")
    add_body_p("When a spam message (e.g. 'Congratulations! You won a free cash prize of ₹50,000. Call to claim now!') is submitted, the system highlights a prominent red warning card with an alert icon, confidence badge, and explanatory notice.")
    add_figure("static/doc_ui_spam_output.png", "Figure 4: Real-Time UI Output - Spam Message Detection with Confidence Score")

    # UI Screenshot 3: Ham Output
    add_heading_2("5.3 Legitimate (Ham) Detection Output Interface")
    add_body_p("When a valid personal message (e.g. 'Hey, are we still meeting tomorrow for lunch at 1 PM?') is entered, the UI renders a green 'Legitimate Message' shield card with positive feedback.")
    add_figure("static/doc_ui_ham_output.png", "Figure 5: Real-Time UI Output - Legitimate (Ham) Message Detection")

    # --- 6. CONCLUSION & FUTURE ENHANCEMENTS ---
    add_heading_1("6. Conclusion & Future Enhancements")
    add_body_p("The Email Spam Classifier project successfully demonstrates a machine learning pipeline integrated with a web web interface. Achieving 100% precision ensures users do not lose critical messages while enjoying reliable protection against unsolicited spam.")

    add_heading_2("Future Enhancements:")
    add_bullet_p("Advanced Transformer Models: ", "Integrate BERT or DistilBERT for deep contextual understanding of complex phishing attacks.")
    add_bullet_p("Browser Extension: ", "Develop a Chrome/Firefox extension to scan emails directly inside webmail clients like Gmail and Outlook.")
    add_bullet_p("RESTful API: ", "Expose an authenticated API endpoint for enterprise integration with mail server gateways.")

    output_filename = "Email_Spam_Classifier_Project_Report.docx"
    doc.save(output_filename)
    print(f"Report saved successfully as '{output_filename}'")

if __name__ == "__main__":
    create_report()
